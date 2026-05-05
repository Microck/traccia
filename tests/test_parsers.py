from __future__ import annotations

import json
import mailbox
import types
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

import traccia.document_normalizer as document_normalizer
from traccia.config import TracciaConfig
from traccia.models import AttachmentKind, SourceCategory, SourceFamily, SourceType
from traccia.parsers import ingestable_file, parse_document


def test_parse_ai_conversation_export_builds_chat_spans_with_stable_line_numbers(
    tmp_path: Path,
) -> None:
    export_path = tmp_path / "assistant-conversation.json"
    export_path.write_text(
        json.dumps(
            {
                "title": "Test Conversation",
                "messages": [
                    {
                        "role": "user",
                        "content": "First prompt",
                        "created_at": "2026-04-19T20:00:00Z",
                    },
                    {
                        "role": "assistant",
                        "content": "Second answer\nwith continuation",
                        "created_at": "2026-04-19T20:01:00Z",
                    },
                ],
            }
        )
    )

    parsed = parse_document(
        export_path,
        project_relative_path=Path("assistant-conversation.json"),
    )

    assert parsed.source.source_type == SourceType.CHAT
    assert parsed.source.source_category == SourceCategory.AI_DIALOGUE
    assert parsed.source.parser == "ai-conversation-json"
    assert len(parsed.spans) == 2
    assert parsed.spans[0].line_start == 1
    assert parsed.spans[0].line_end == 1
    assert parsed.spans[1].line_start == 3
    assert parsed.spans[1].line_end == 4
    assert parsed.text == (
        "User: First prompt\n\n"
        "Assistant: Second answer\nwith continuation"
    )


def test_parse_instagram_export_html_uses_family_normalizer(tmp_path: Path) -> None:
    export_path = tmp_path / "message_1.html"
    export_path.write_text(
        """
        <html>
          <head><style>body { color: red; }</style></head>
          <body>
            <div class="thread">
              <h1>h</h1>
              <div>no way you know them?</div>
              <a href="https://www.instagram.com/stories/example">story link</a>
              <div>may. 13, 2024 9:16 pm</div>
            </div>
          </body>
        </html>
        """
    )

    parsed = parse_document(
        export_path,
        project_relative_path=Path(
            "Instagram/export/your_instagram_activity/messages/inbox/h_123/message_1.html"
        ),
        source_family=SourceFamily.INSTAGRAM_EXPORT,
        source_family_subproduct="messages",
    )

    assert parsed.source.parser == "instagram_export-html"
    assert parsed.source.metadata["family_normalizer"] == "html-export"
    assert parsed.source.metadata["family_normalizer_record_count"] >= 1
    assert "body { color: red; }" not in parsed.text
    assert "no way you know them?" in parsed.text
    assert "https://www.instagram.com/stories/example" in parsed.text


def test_parse_twitter_archive_js_uses_family_normalizer(tmp_path: Path) -> None:
    export_path = tmp_path / "account.js"
    export_path.write_text(
        """
        window.YTD.account.part0 = [
          {
            "account": {
              "username": "JustMicrock",
              "email": "archive-owner@example.com",
              "createdAt": "2019-12-02T14:21:19.009Z"
            }
          }
        ];
        """
    )

    parsed = parse_document(
        export_path,
        project_relative_path=Path("Twitter/archive/data/account.js"),
        source_family=SourceFamily.TWITTER_ARCHIVE,
        source_family_subproduct="account",
    )

    assert parsed.source.parser == "twitter-ytd-json"
    assert parsed.source.metadata["family_normalizer"] == "twitter-ytd-js"
    assert parsed.source.metadata["family_normalizer_kind"] == "account"
    assert "username: JustMicrock" in parsed.text
    assert "email: archive-owner@example.com" in parsed.text


def test_parse_remote_media_url_skips_enrichment_when_summarize_missing(tmp_path: Path) -> None:
    note_path = tmp_path / "notes.md"
    note_path.write_text("watch later https://youtu.be/example and revisit the packaging flow")
    config = TracciaConfig.model_validate(
        {
            "multimodal": {
                "remote_media_enrichment_command": str(tmp_path / "missing-summarize"),
            }
        }
    )

    parsed = parse_document(
        note_path,
        project_relative_path=Path("exports/notes.md"),
        config=config,
    )

    assert parsed.source.metadata["remote_media_reference_count"] == 1
    assert len(parsed.attachments) == 1
    assert parsed.attachments[0].reference == "https://youtu.be/example"
    assert parsed.attachments[0].extracted_text is None


def test_parse_reddit_export_csv_uses_family_normalizer(tmp_path: Path) -> None:
    export_path = tmp_path / "comments.csv"
    export_path.write_text(
        "subreddit,body,created_utc\nLocalLLaMA,benchmarked OCR tools,1715577360\n"
    )

    parsed = parse_document(
        export_path,
        project_relative_path=Path("Reddit/reddit-export/comments.csv"),
        source_family=SourceFamily.REDDIT_EXPORT,
        source_family_subproduct="comments",
    )

    assert parsed.source.parser == "reddit-csv"
    assert parsed.source.metadata["family_normalizer"] == "reddit-csv"
    assert parsed.source.metadata["family_normalizer_record_count"] == 1
    assert "subreddit: LocalLLaMA" in parsed.text
    assert "benchmarked OCR tools" in parsed.text


def test_parse_google_takeout_youtube_csv_uses_family_normalizer(tmp_path: Path) -> None:
    export_path = tmp_path / "suscripciones.csv"
    export_path.write_text(
        (
            "ID del canal,URL del canal,Título del canal\n"
            "UC123,http://www.youtube.com/channel/UC123,Disrupt\n"
        ),
        encoding="utf-8",
    )

    parsed = parse_document(
        export_path,
        project_relative_path=Path("Takeout/YouTube y YouTube Music/suscripciones/suscripciones.csv"),
        source_family=SourceFamily.GOOGLE_TAKEOUT,
        source_family_subproduct="youtube-and-youtube-music",
    )

    assert parsed.source.parser == "google-takeout-csv"
    assert parsed.source.source_category == SourceCategory.PLATFORM_EXPORT_ACTIVITY
    assert parsed.source.metadata["family_normalizer"] == "google-takeout-csv"
    assert parsed.source.metadata["family_normalizer_record_count"] == 1
    assert "channel: Disrupt" in parsed.text
    assert len(parsed.attachments) == 1
    assert parsed.attachments[0].reference == "http://www.youtube.com/channel/UC123"


def test_parse_google_takeout_profile_json_uses_family_normalizer(tmp_path: Path) -> None:
    export_path = tmp_path / "Perfil.json"
    export_path.write_text(
        json.dumps(
            {
                "name": {
                    "givenName": "Microck",
                    "formattedName": "Microck",
                },
                "displayName": "Microck",
                "emails": [{"value": "profile-owner@example.com"}],
            }
        ),
        encoding="utf-8",
    )

    parsed = parse_document(
        export_path,
        project_relative_path=Path("Takeout/Perfil/Perfil.json"),
        source_family=SourceFamily.GOOGLE_TAKEOUT,
        source_family_subproduct="profile",
    )

    assert parsed.source.parser == "google-takeout-json"
    assert parsed.source.metadata["family_normalizer"] == "google-takeout-json"
    assert "displayName: Microck" in parsed.text
    assert "emails.value: profile-owner@example.com" in parsed.text


def test_parse_google_takeout_calendar_ics_uses_family_normalizer(tmp_path: Path) -> None:
    export_path = tmp_path / "calendar.ics"
    export_path.write_text(
        "\n".join(
            [
                "BEGIN:VCALENDAR",
                "BEGIN:VEVENT",
                "DTSTART:20230408T100000Z",
                "DTEND:20230408T130000Z",
                "SUMMARY:Comida con Jimbo",
                "DESCRIPTION:Lunch",
                "END:VEVENT",
                "END:VCALENDAR",
            ]
        ),
        encoding="utf-8",
    )

    parsed = parse_document(
        export_path,
        project_relative_path=Path("Takeout/Calendar/calendar-owner@example.com.ics"),
        source_family=SourceFamily.GOOGLE_TAKEOUT,
        source_family_subproduct="calendar",
    )

    assert parsed.source.source_type == SourceType.CALENDAR
    assert parsed.source.parser == "google-takeout-calendar"
    assert parsed.source.metadata["family_normalizer"] == "google-takeout-calendar"
    assert "summary: Comida con Jimbo" in parsed.text


def test_parse_google_takeout_mbox_uses_family_normalizer(tmp_path: Path) -> None:
    export_path = tmp_path / "all-mail.mbox"
    mbox = mailbox.mbox(export_path)
    message = mailbox.mboxMessage()
    message["From"] = "microck@example.com"
    message["To"] = "friend@example.com"
    message["Date"] = "Mon, 10 Nov 2025 17:01:52 +0100"
    message["Subject"] = "CCPA DELETE"
    message.set_payload("Please delete my data.")
    mbox.add(message)
    mbox.flush()

    parsed = parse_document(
        export_path,
        project_relative_path=Path("Takeout/Correo/Todo el correo, incluido Spam y Papelera.mbox"),
        source_family=SourceFamily.GOOGLE_TAKEOUT,
        source_family_subproduct="mail",
    )

    assert parsed.source.parser == "google-takeout-mbox"
    assert parsed.source.metadata["family_normalizer"] == "google-takeout-mbox"
    assert parsed.source.metadata["family_normalizer_record_count"] == 1
    assert "subject: CCPA DELETE" in parsed.text
    assert "Please delete my data." in parsed.text


def test_ingestable_file_skips_known_binary_media_extensions(tmp_path: Path) -> None:
    image_path = tmp_path / "photo.webp"
    image_path.write_bytes(b"RIFFxxxxWEBPVP8 " + (b"\x00" * 32))

    assert ingestable_file(image_path) is False


def test_ingestable_file_skips_large_extensionless_files_without_content_sniff(
    tmp_path: Path,
    monkeypatch,
) -> None:
    media_path = tmp_path / "1913559960169918466-1913559960169918466"
    with media_path.open("wb") as handle:
        handle.seek(1024 * 1024 + 1)
        handle.write(b"\x00")

    original_open = Path.open

    def fail_if_media_opened(self: Path, *args, **kwargs):
        if self == media_path:
            raise AssertionError("large extensionless files should not be opened for text sniffing")
        return original_open(self, *args, **kwargs)

    monkeypatch.setattr(Path, "open", fail_if_media_opened)

    assert ingestable_file(media_path) is False


def test_ingestable_file_accepts_google_ai_studio_extensionless_exports_without_sniff(
    tmp_path: Path,
    monkeypatch,
) -> None:
    export_path = tmp_path / "Takeout" / "Drive" / "Google AI Studio" / "Configurar Disco"
    export_path.parent.mkdir(parents=True)
    export_path.write_text("I configured disk partitions.\n")

    original_open = Path.open

    def fail_if_export_opened(self: Path, *args, **kwargs):
        if self == export_path:
            raise AssertionError("extensionless Google AI Studio exports should not be sniffed")
        return original_open(self, *args, **kwargs)

    monkeypatch.setattr(Path, "open", fail_if_export_opened)

    assert ingestable_file(export_path) is True


def test_parse_docx_uses_native_document_normalizer_when_requested(tmp_path: Path) -> None:
    document_path = tmp_path / "notes.docx"
    document = DocxDocument()
    document.add_paragraph("I built a parser.")
    document.save(document_path)
    config = TracciaConfig.model_validate(
        {
            "document_normalization": {
                "provider": "native",
            }
        }
    )

    parsed = parse_document(
        document_path,
        project_relative_path=Path("notes.docx"),
        config=config,
    )

    assert parsed.source.source_type == SourceType.DOCX
    assert parsed.source.parser == "python-docx"
    assert parsed.source.metadata["document_normalizer"] == "native"
    assert parsed.source.metadata["document_ocr_used"] is False
    assert "I built a parser." in parsed.text


def test_parse_xlsx_workbook_as_authored_spreadsheet(tmp_path: Path) -> None:
    document_path = tmp_path / "checklist.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Build"
    worksheet.append(["Task", "Status"])
    worksheet.append(["Flash keyboard firmware", "done"])
    worksheet.append(["Tune stabilizers", "done"])
    workbook.save(document_path)

    parsed = parse_document(
        document_path,
        project_relative_path=Path("Takeout/Drive/checklist.xlsx"),
    )

    assert parsed.source.source_type == SourceType.SPREADSHEET
    assert parsed.source.source_category == SourceCategory.AUTHORED_CONTENT
    assert parsed.source.parser == "xlsx"
    assert "# Sheet: Build" in parsed.text
    assert "Flash keyboard firmware | done" in parsed.text


def test_parse_docx_auto_prefers_docling_when_available(tmp_path: Path, monkeypatch) -> None:
    document_path = tmp_path / "notes.docx"
    document_path.write_bytes(b"fake-docx")
    captured: dict[str, object] = {}

    class FakeDocumentConverter:
        def __init__(self, *args, **kwargs) -> None:
            captured["args"] = args
            captured["kwargs"] = kwargs

        def convert(self, source: Path):
            captured["source"] = source
            return types.SimpleNamespace(
                document=types.SimpleNamespace(
                    export_to_markdown=lambda: "# Notes\n\nI built a parser."
                )
            )

    def fake_import_module(name: str):
        if name == "docling.document_converter":
            return types.SimpleNamespace(DocumentConverter=FakeDocumentConverter)
        if name == "docling.datamodel.base_models":
            return types.SimpleNamespace(InputFormat=types.SimpleNamespace(PDF="pdf"))
        if name == "docling.datamodel.pipeline_options":
            return types.SimpleNamespace(PdfPipelineOptions=lambda: None)
        raise ModuleNotFoundError(name)

    monkeypatch.setattr(document_normalizer.importlib, "import_module", fake_import_module)

    parsed = parse_document(
        document_path,
        project_relative_path=Path("notes.docx"),
    )

    assert parsed.source.parser == "docling"
    assert parsed.source.metadata["document_normalizer"] == "docling"
    assert parsed.source.metadata["document_normalization_mode"] == "markdown"
    assert parsed.source.metadata["document_ocr_provider"] == "none"
    assert captured["source"] == document_path
    assert parsed.text.startswith("# Notes")


def test_parse_pdf_auto_prefers_marker_when_available(tmp_path: Path, monkeypatch) -> None:
    document_path = tmp_path / "paper.pdf"
    document_path.write_bytes(b"%PDF-1.4 fake")
    captured: dict[str, object] = {}

    def fake_run(command, *, check, capture_output, text):
        del check, capture_output, text
        captured["command"] = command
        output_dir = Path(command[command.index("--output_dir") + 1])
        output_dir.mkdir(parents=True, exist_ok=True)
        (output_dir / "paper").mkdir(exist_ok=True)
        (output_dir / "paper" / "paper.md").write_text("# Parsed\n\nMarker output")
        return types.SimpleNamespace(returncode=0, stdout="", stderr="")

    monkeypatch.setattr(document_normalizer.subprocess, "run", fake_run)

    parsed = parse_document(
        document_path,
        project_relative_path=Path("paper.pdf"),
    )

    assert parsed.source.parser == "marker"
    assert parsed.source.metadata["document_normalizer"] == "marker"
    assert parsed.source.metadata["document_ocr_provider"] == "marker-force-ocr"
    assert "--force_ocr" in captured["command"]
    assert "Marker output" in parsed.text


def test_parse_docx_auto_falls_back_to_markitdown_when_docling_missing(
    tmp_path: Path,
    monkeypatch,
) -> None:
    document_path = tmp_path / "notes.docx"
    document_path.write_bytes(b"fake-docx")
    captured: dict[str, object] = {}

    def fake_import_module(name: str):
        if name == "docling.document_converter":
            raise ModuleNotFoundError(name)
        if name == "markitdown":
            class FakeMarkItDown:
                def __init__(self, *, enable_plugins: bool = False, **kwargs) -> None:
                    captured["enable_plugins"] = enable_plugins
                    captured["kwargs"] = kwargs

                def convert(self, source: Path):
                    captured["source"] = source
                    return types.SimpleNamespace(text_content="# Notes\n\nFallback parser.")

            return types.SimpleNamespace(MarkItDown=FakeMarkItDown)
        raise ModuleNotFoundError(name)

    monkeypatch.setattr(document_normalizer.importlib, "import_module", fake_import_module)

    parsed = parse_document(
        document_path,
        project_relative_path=Path("notes.docx"),
    )

    assert parsed.source.parser == "markitdown"
    assert parsed.source.metadata["document_normalizer"] == "markitdown"
    assert parsed.source.metadata["document_ocr_provider"] == "none"
    assert captured["enable_plugins"] is False
    assert captured["source"] == document_path


def test_parse_pdf_auto_falls_back_to_docling_when_marker_missing(tmp_path: Path, monkeypatch) -> None:
    document_path = tmp_path / "scan.pdf"
    document_path.write_bytes(b"%PDF-1.4 fake")
    captured: dict[str, object] = {}

    def fake_run(command, *, check, capture_output, text):
        del command, check, capture_output, text
        raise FileNotFoundError("marker_single")

    class FakePdfPipelineOptions:
        def __init__(self) -> None:
            self.do_ocr = False
            self.ocr_options = None

    class FakeOcrAutoOptions:
        pass

    class FakeDocumentConverter:
        def __init__(self, *, format_options=None) -> None:
            captured["format_options"] = format_options

        def convert(self, source: Path):
            captured["source"] = source
            return types.SimpleNamespace(
                document=types.SimpleNamespace(export_to_markdown=lambda: "Docling fallback")
            )

    class FakePdfFormatOption:
        def __init__(self, *, pipeline_options) -> None:
            captured["pipeline_options"] = pipeline_options

    def fake_import_module(name: str):
        if name == "docling.document_converter":
            return types.SimpleNamespace(
                DocumentConverter=FakeDocumentConverter,
                PdfFormatOption=FakePdfFormatOption,
            )
        if name == "docling.datamodel.base_models":
            return types.SimpleNamespace(InputFormat=types.SimpleNamespace(PDF="pdf"))
        if name == "docling.datamodel.pipeline_options":
            return types.SimpleNamespace(
                PdfPipelineOptions=FakePdfPipelineOptions,
                OcrAutoOptions=FakeOcrAutoOptions,
            )
        raise ModuleNotFoundError(name)

    monkeypatch.setattr(document_normalizer.subprocess, "run", fake_run)
    monkeypatch.setattr(document_normalizer.importlib, "import_module", fake_import_module)

    parsed = parse_document(
        document_path,
        project_relative_path=Path("scan.pdf"),
    )

    assert parsed.source.parser == "docling"
    assert parsed.source.metadata["document_normalizer"] == "docling"
    assert parsed.source.metadata["document_ocr_provider"] == "docling-auto"
    assert parsed.text == "Docling fallback"


def test_parse_docx_markitdown_missing_dependency_raises_clear_error(
    tmp_path: Path,
    monkeypatch,
) -> None:
    document_path = tmp_path / "notes.docx"
    document_path.write_bytes(b"fake-docx")

    def fake_import_module(name: str):
        raise ModuleNotFoundError(name)

    monkeypatch.setattr(document_normalizer.importlib, "import_module", fake_import_module)
    config = TracciaConfig.model_validate(
        {
            "document_normalization": {
                "provider": "markitdown",
            },
        }
    )

    with pytest.raises(RuntimeError, match="markitdown is not installed"):
        parse_document(
            document_path,
            project_relative_path=Path("notes.docx"),
            config=config,
        )


def test_parse_pdf_docling_uses_local_ocr_backend(
    tmp_path: Path,
    monkeypatch,
) -> None:
    document_path = tmp_path / "scan.pdf"
    document_path.write_bytes(b"%PDF-1.4 fake")
    captured: dict[str, object] = {}

    class FakePdfPipelineOptions:
        def __init__(self) -> None:
            self.do_ocr = False
            self.ocr_options = None

    class FakeRapidOcrOptions:
        def __init__(self) -> None:
            captured["ocr_options_class"] = "RapidOcrOptions"

    class FakeDocumentConverter:
        def __init__(self, *, format_options=None) -> None:
            captured["format_options"] = format_options

        def convert(self, source: Path):
            captured["source"] = source
            return types.SimpleNamespace(
                document=types.SimpleNamespace(
                    export_to_markdown=lambda: "Recovered local OCR markdown"
                )
            )

    class FakePdfFormatOption:
        def __init__(self, *, pipeline_options) -> None:
            captured["pipeline_options"] = pipeline_options

    def fake_import_module(name: str):
        if name == "docling.document_converter":
            return types.SimpleNamespace(
                DocumentConverter=FakeDocumentConverter,
                PdfFormatOption=FakePdfFormatOption,
            )
        if name == "docling.datamodel.base_models":
            return types.SimpleNamespace(InputFormat=types.SimpleNamespace(PDF="pdf"))
        if name == "docling.datamodel.pipeline_options":
            return types.SimpleNamespace(
                PdfPipelineOptions=FakePdfPipelineOptions,
                RapidOcrOptions=FakeRapidOcrOptions,
            )
        raise ModuleNotFoundError(name)

    monkeypatch.setattr(document_normalizer.importlib, "import_module", fake_import_module)
    config = TracciaConfig.model_validate(
        {
            "document_normalization": {
                "provider": "docling",
                "ocr_provider": "rapidocr",
            },
        }
    )

    parsed = parse_document(
        document_path,
        project_relative_path=Path("scan.pdf"),
        config=config,
    )

    assert parsed.source.source_type == SourceType.PDF
    assert parsed.source.parser == "docling"
    assert parsed.source.metadata["document_normalizer"] == "docling"
    assert parsed.source.metadata["document_ocr_provider"] == "docling-rapidocr"
    assert parsed.source.metadata["document_ocr_requested"] is True
    assert parsed.source.metadata["document_ocr_used"] is True
    assert captured["ocr_options_class"] == "RapidOcrOptions"
    assert captured["pipeline_options"].do_ocr is True
    assert isinstance(captured["pipeline_options"].ocr_options, FakeRapidOcrOptions)
    assert captured["source"] == document_path


def test_parse_pdf_marker_can_be_disabled_explicitly(tmp_path: Path, monkeypatch) -> None:
    document_path = tmp_path / "paper.pdf"
    document_path.write_bytes(b"%PDF-1.4 fake")
    captured: dict[str, object] = {}

    class FakePdfPipelineOptions:
        def __init__(self) -> None:
            self.do_ocr = False
            self.ocr_options = None

    class FakeDocumentConverter:
        def __init__(self, *, format_options=None) -> None:
            captured["format_options"] = format_options

        def convert(self, source: Path):
            captured["source"] = source
            return types.SimpleNamespace(
                document=types.SimpleNamespace(export_to_markdown=lambda: "Docling only")
            )

    class FakePdfFormatOption:
        def __init__(self, *, pipeline_options) -> None:
            captured["pipeline_options"] = pipeline_options

    def fake_import_module(name: str):
        if name == "docling.document_converter":
            return types.SimpleNamespace(
                DocumentConverter=FakeDocumentConverter,
                PdfFormatOption=FakePdfFormatOption,
            )
        if name == "docling.datamodel.base_models":
            return types.SimpleNamespace(InputFormat=types.SimpleNamespace(PDF="pdf"))
        if name == "docling.datamodel.pipeline_options":
            return types.SimpleNamespace(
                PdfPipelineOptions=FakePdfPipelineOptions,
                OcrAutoOptions=lambda: object(),
            )
        raise ModuleNotFoundError(name)

    def forbidden_run(command, *, check, capture_output, text):
        raise AssertionError(f"marker should not run: {command}")

    monkeypatch.setattr(document_normalizer.importlib, "import_module", fake_import_module)
    monkeypatch.setattr(document_normalizer.subprocess, "run", forbidden_run)
    config = TracciaConfig.model_validate(
        {
            "document_normalization": {
                "provider": "docling",
                "ocr_provider": "none",
            },
        }
    )

    parsed = parse_document(
        document_path,
        project_relative_path=Path("paper.pdf"),
        config=config,
    )

    assert parsed.source.parser == "docling"
    assert parsed.source.metadata["document_ocr_provider"] == "none"
    assert parsed.text == "Docling only"

def test_parse_html_attachment_transcribes_local_video_when_tools_are_available(
    tmp_path: Path,
    monkeypatch,
) -> None:
    export_path = tmp_path / "post.html"
    video_path = tmp_path / "clip.mp4"
    export_path.write_text(
        """
        <html>
          <body>
            <p>I posted a build update.</p>
            <video src="clip.mp4"></video>
          </body>
        </html>
        """
    )
    video_path.write_bytes(b"fake-video")

    def fake_which(command: str) -> str | None:
        if command in {"ffmpeg", "whisper"}:
            return f"/usr/bin/{command}"
        return None

    def fake_run(command, *, capture_output, text, timeout, check):
        del capture_output, text, timeout, check
        if command[0] == "ffmpeg":
            Path(command[-1]).write_bytes(b"RIFFfake-audio")
            return types.SimpleNamespace(returncode=0, stdout="", stderr="")
        if command[0] == "whisper":
            audio_path = Path(command[1])
            output_dir = Path(command[command.index("--output_dir") + 1])
            (output_dir / f"{audio_path.stem}.json").write_text(
                json.dumps({"text": "I built the parser and shipped the fix."})
            )
            return types.SimpleNamespace(returncode=0, stdout="", stderr="")
        raise AssertionError(f"unexpected command: {command}")

    monkeypatch.setattr("traccia.parsers.shutil.which", fake_which)
    monkeypatch.setattr("traccia.parsers.subprocess.run", fake_run)

    parsed = parse_document(
        export_path,
        project_relative_path=Path("exports/post.html"),
    )

    assert len(parsed.attachments) == 1
    attachment = parsed.attachments[0]
    assert attachment.kind == AttachmentKind.VIDEO
    assert attachment.extracted_text == "I built the parser and shipped the fix."
    assert attachment.metadata["transcription_provider"] == "whisper_cli"
    assert attachment.metadata["transcription_model"] == "turbo"

def test_parse_html_attachment_can_disable_local_media_transcription(
    tmp_path: Path,
) -> None:
    export_path = tmp_path / "post.html"
    video_path = tmp_path / "clip.mp4"
    export_path.write_text(
        """
        <html>
          <body>
            <video src="clip.mp4"></video>
          </body>
        </html>
        """
    )
    video_path.write_bytes(b"fake-video")
    config = TracciaConfig.model_validate(
        {
            "multimodal": {
                "enable_local_media_transcription": False,
            }
        }
    )

    parsed = parse_document(
        export_path,
        project_relative_path=Path("exports/post.html"),
        config=config,
    )

    assert len(parsed.attachments) == 1
    assert parsed.attachments[0].extracted_text is None
